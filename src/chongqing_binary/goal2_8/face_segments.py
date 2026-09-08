"""Segment the Face task session into its emotion-movie phases.

`面部2-任务` videos are 11-minute multi-phase sessions, not single clips. Goal
2.7 sampled 16 frames uniformly across the whole session, which mixes film
watching, reading aloud and interview, and mostly encodes identity, appearance
and room background. Segment boundaries come from `附件/网页数据.xlsx`, aligned
to the recording by the `start.mp3` landmark.
"""

from __future__ import annotations

import json
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from ..paradigm import load_paradigm_spec
from ..paradigm.spec import ParadigmSpec
from .config import cv_subjects, ensure_output, load_goal_config, project_path

AUDIO_SR = 8000


def _to_seconds(value: Any) -> float | None:
    """Web-log timestamps are H:M:S:ms."""
    if value is None:
        return None
    parts = str(value).strip().split(":")
    if len(parts) != 4:
        return None
    try:
        hours, minutes, seconds, millis = (int(p) for p in parts)
    except ValueError:
        return None
    return hours * 3600 + minutes * 60 + seconds + millis / 1000.0


def read_web_log(path: str | Path) -> tuple[list[str], list[dict[str, Any]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(iterator)]
    records = []
    for row in iterator:
        if any(v is not None for v in row):
            records.append({h: row[i] for i, h in enumerate(header) if h})
    workbook.close()
    return header, records


def _extract_audio(video: Path, out_wav: Path, duration: float) -> bool:
    result = subprocess.run(
        ["ffmpeg", "-y", "-t", str(duration), "-i", str(video), "-vn", "-acodec", "pcm_s16le",
         "-ar", str(AUDIO_SR), "-ac", "1", "-loglevel", "error", str(out_wav)],
        capture_output=True,
    )
    return result.returncode == 0 and out_wav.exists()


def _normalised_correlation(signal: np.ndarray, template: np.ndarray) -> np.ndarray:
    from scipy.signal import fftconvolve

    n = len(template)
    energy = fftconvolve(signal ** 2, np.ones(n), mode="valid")
    cross = fftconvolve(signal, template[::-1], mode="valid")
    return cross / np.sqrt(np.maximum(energy * np.sum(template ** 2), 1e-9))


def detect_landmark(video: Path, template: np.ndarray, search_sec: float,
                    scratch: Path) -> tuple[float, float]:
    """Return (offset_seconds, correlation) of the best landmark match."""
    from scipy.io import wavfile

    wav = scratch / f"{video.stem}.wav"
    try:
        if not _extract_audio(video, wav, search_sec):
            return float("nan"), 0.0
        _, audio = wavfile.read(wav)
        audio = np.asarray(audio, dtype=float)
        if audio.ndim > 1:
            audio = audio.mean(axis=1)
        if audio.size <= template.size:
            return float("nan"), 0.0
        corr = _normalised_correlation(audio, template)
        best = int(np.argmax(corr))
        return best / AUDIO_SR, float(corr[best])
    except Exception:  # noqa: BLE001
        return float("nan"), 0.0
    finally:
        wav.unlink(missing_ok=True)


def load_landmark_template(scratch: Path, spec: ParadigmSpec) -> np.ndarray:
    """Extract the start.mp3 landmark from the stimulus bundle."""
    import zipfile

    from scipy.io import wavfile

    raw_root = project_path(load_goal_config("configs/goal2_8/face.yaml")["paths"]["raw_data_dir"])
    bundle = raw_root / spec.face["task"]["alignment_landmark"].split("::")[0]
    member = spec.face["task"]["alignment_landmark"].split("::")[1]
    scratch.mkdir(parents=True, exist_ok=True)
    mp3 = scratch / "landmark.mp3"
    with zipfile.ZipFile(bundle) as archive:
        name = next(n for n in archive.namelist() if n.endswith(member.split("/")[-1]))
        mp3.write_bytes(archive.read(name))
    wav = scratch / "landmark.wav"
    if not _extract_audio(mp3, wav, 10.0):
        raise RuntimeError("could not decode the alignment landmark")
    _, template = wavfile.read(wav)
    return np.asarray(template, dtype=float)


def _video_duration(path: Path) -> float:
    result = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(path)],
        capture_output=True, text=True,
    )
    try:
        return float(result.stdout.strip())
    except ValueError:
        return float("nan")


@dataclass
class SubjectSegments:
    l_id: str
    a_id: str
    status: str
    reason: str = ""
    lead_in_sec: float = float("nan")
    landmark_corr: float = 0.0
    landmark_source: str = ""
    video_duration_sec: float = float("nan")
    segments: dict[str, tuple[float, float]] | None = None
    duration_mismatch: str = ""


def build_segment_index(config_path: str | Path = "configs/goal2_8/face.yaml",
                        limit: int | None = None, n_jobs: int | None = None,
                        spec: ParadigmSpec | None = None) -> dict[str, Any]:
    from joblib import Parallel, delayed

    config = load_goal_config(config_path)
    spec = spec or load_paradigm_spec(config["paths"]["paradigm_spec"])
    face_cfg = config["face"]
    task_face = spec.face["task"]
    movies = spec.face_movies()

    raw_root = project_path(config["paths"]["raw_data_dir"])
    video_dir = raw_root / task_face["dir"]
    _, records = read_web_log(raw_root / task_face["segment_source"])
    by_aid = {str(r[task_face["key_column"]]).strip(): r for r in records if r.get(task_face["key_column"])}

    split = cv_subjects(config)
    scratch = Path(ensure_output(f"{face_cfg['outputs_dir']}/_scratch/.keep", config)).parent
    template = load_landmark_template(scratch, spec)

    targets: list[tuple[str, str, Path]] = []
    for row in split.itertuples():
        l_id, a_id = str(row.L_id), str(row.A_id)
        video = video_dir / f"{l_id}.mp4"
        if a_id in by_aid and video.exists():
            targets.append((l_id, a_id, video))
    if limit:
        targets = targets[:limit]

    search_sec = float(face_cfg["landmark_search_sec"])
    min_corr = float(face_cfg["landmark_min_correlation"])
    tolerance = float(task_face["segment_duration_tolerance_sec"])

    def work(l_id: str, a_id: str, video: Path) -> SubjectSegments:
        record = by_aid[a_id]
        anchor = _to_seconds(record.get(movies[task_face["movie_order"][0]]["start_column"]))
        if anchor is None:
            return SubjectSegments(l_id, a_id, "blocked", "missing_anchor_timestamp")
        offset, corr = detect_landmark(video, template, search_sec, scratch)
        duration = _video_duration(video)
        segments: dict[str, tuple[float, float]] = {}
        mismatches: list[str] = []
        for name in task_face["movie_order"]:
            cfg = movies[name]
            start = _to_seconds(record.get(cfg["start_column"]))
            end = _to_seconds(record.get(cfg["end_column"]))
            if start is None or end is None or end <= start:
                mismatches.append(f"{name}:missing")
                continue
            expected = float(cfg["stimulus_duration_sec"])
            if abs((end - start) - expected) > tolerance:
                mismatches.append(f"{name}:duration")
                continue
            segments[name] = (start - anchor, end - anchor)   # relative to the anchor
        return SubjectSegments(
            l_id, a_id, "ok" if segments else "blocked",
            "" if segments else "no_valid_segments",
            lead_in_sec=offset, landmark_corr=corr,
            landmark_source="landmark" if corr >= min_corr else "cohort_median",
            video_duration_sec=duration, segments=segments,
            duration_mismatch=";".join(mismatches),
        )

    jobs = int(n_jobs if n_jobs is not None else config.get("run", {}).get("n_jobs", 1))
    if jobs == 1:
        results = [work(*t) for t in targets]
    else:
        results = Parallel(n_jobs=jobs, backend="threading", verbose=5)(delayed(work)(*t) for t in targets)

    confident = [r.lead_in_sec for r in results if r.landmark_corr >= min_corr and np.isfinite(r.lead_in_sec)]
    fallback = float(np.median(confident)) if confident else float(face_cfg["landmark_default_lead_in_sec"])

    rows: list[dict[str, Any]] = []
    for r in results:
        lead = r.lead_in_sec if r.landmark_corr >= min_corr and np.isfinite(r.lead_in_sec) else fallback
        row: dict[str, Any] = {
            "L_id": r.l_id, "A_id": r.a_id, "status": r.status, "reason": r.reason,
            "lead_in_sec": lead, "landmark_corr": r.landmark_corr,
            "lead_in_source": r.landmark_source, "video_duration_sec": r.video_duration_sec,
            "duration_mismatch": r.duration_mismatch,
        }
        for name in task_face["movie_order"]:
            seg = (r.segments or {}).get(name)
            if seg:
                start, end = seg[0] + lead, seg[1] + lead
                if np.isfinite(r.video_duration_sec) and end > r.video_duration_sec:
                    row[f"{name}_start_sec"] = np.nan
                    row[f"{name}_end_sec"] = np.nan
                    row["reason"] = (row["reason"] + ";" if row["reason"] else "") + f"{name}:beyond_video_end"
                else:
                    row[f"{name}_start_sec"] = start
                    row[f"{name}_end_sec"] = end
            else:
                row[f"{name}_start_sec"] = np.nan
                row[f"{name}_end_sec"] = np.nan
        rows.append(row)

    index = pd.DataFrame(rows)
    out_dir = Path(ensure_output(f"{face_cfg['outputs_dir']}/.keep", config)).parent
    index_path = out_dir / "segment_index.csv"
    index.to_csv(index_path, index=False)

    usable = index[[f"{n}_start_sec" for n in task_face["movie_order"]]].notna().all(axis=1)
    manifest = {
        "subjects_attempted": len(targets),
        "subjects_indexed": int(len(index)),
        "subjects_all_three_movies": int(usable.sum()),
        "landmark_confident": int((index["lead_in_source"] == "landmark").sum()),
        "landmark_median_lead_in_sec": fallback,
        "movie_order": task_face["movie_order"],
        "block_reasons": index.loc[index["status"] != "ok", "reason"].value_counts().to_dict(),
        "duration_mismatch_subjects": int((index["duration_mismatch"] != "").sum()),
        "segment_index": str(index_path),
    }
    (out_dir / "segment_index_manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    return manifest
