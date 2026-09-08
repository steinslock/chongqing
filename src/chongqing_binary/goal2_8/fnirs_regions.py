"""Channel-to-region mapping for both fNIRS devices.

Goal 2.7 recorded `qc_region_mapping_status: unconfirmed_channel_global_hemisphere_only`.
Both devices in fact ship a channel-to-cortex mapping in the attachments:
Yiruid via MNI projections of the yrd-53 montage, Bikom via a network partition.
"""

from __future__ import annotations

import csv
from functools import lru_cache
from pathlib import Path
from typing import Any

from ..paths import raw_data_root


def _laterality(x: float, midline_mm: float = 20.0) -> str:
    if x < -midline_mm:
        return "left"
    if x > midline_mm:
        return "right"
    return "mid"


def _anterior_posterior(y: float) -> str:
    if y > 50.0:
        return "frontopolar"
    if y > 20.0:
        return "anterior"
    if y > -10.0:
        return "central"
    return "posterior"


@lru_cache(maxsize=4)
def yiruid_regions(mni_file: str | None = None) -> dict[int, dict[str, Any]]:
    """CH number -> region labels and MNI coordinates for the yrd-53 montage."""
    path = Path(mni_file) if mni_file else raw_data_root() / "附件/前额叶_20260106093254/前额叶_皮层MNI.csv"
    out: dict[int, dict[str, Any]] = {}
    with Path(path).open("r", encoding="utf-8-sig") as handle:
        for row in csv.reader(handle):
            if not row or not row[0].upper().startswith("CH"):
                continue
            try:
                channel = int(row[0][2:])
                x, y, z = (float(v) for v in row[1:4])
            except (ValueError, IndexError):
                continue
            lateral = _laterality(x)
            anterior = _anterior_posterior(y)
            out[channel] = {
                "mni": (x, y, z),
                "laterality": lateral,
                "anterior_posterior": anterior,
                "region": f"{anterior}_{lateral}",
            }
    return out


@lru_cache(maxsize=4)
def bikom_regions(fc_matrix_file: str | None = None) -> dict[int, dict[str, Any]]:
    """CH number -> network label and MNI coordinates from the provider's partition."""
    import openpyxl

    path = Path(fc_matrix_file) if fc_matrix_file else raw_data_root() / "附件/FC_Matrix_Inf_3D.xlsx"
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    header = [str(h) if h is not None else "" for h in rows[0]]
    ch_i, net_i = header.index("ChID"), header.index("Sub_Net")
    x_i, y_i, z_i = header.index("MNI_x"), header.index("MNI_y"), header.index("MNI_z")

    out: dict[int, dict[str, Any]] = {}
    for row in rows[1:]:
        if row[ch_i] is None:
            continue
        channel = int(row[ch_i])
        x, y, z = (float(row[i]) for i in (x_i, y_i, z_i))
        lateral = _laterality(x)
        anterior = _anterior_posterior(y)
        out[channel] = {
            "mni": (x, y, z),
            "network": str(row[net_i]),
            "laterality": lateral,
            "anterior_posterior": anterior,
            # Anatomical region, defined the same way as Yiruid so the two
            # devices get comparable regional resolution. The provider's
            # four-network partition is kept alongside it.
            "region": f"{anterior}_{lateral}",
        }
    return out


def region_groups(device: str) -> dict[str, list[int]]:
    """Region label -> 1-based channel numbers, for the given device.

    Both devices use the same MNI-derived anatomical scheme so their regional
    resolution is comparable. Bikom additionally keeps the provider's network
    partition, prefixed `net_`.
    """
    mapping = yiruid_regions() if device == "yiruid" else bikom_regions()
    groups: dict[str, list[int]] = {}
    for channel, meta in sorted(mapping.items()):
        groups.setdefault(str(meta["region"]), []).append(channel)
        network = meta.get("network")
        if network:
            groups.setdefault(f"net_{network}", []).append(channel)
    return groups
