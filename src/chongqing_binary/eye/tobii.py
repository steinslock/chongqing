"""Reader for the Tobii Pro Lab xlsx export.

The export is one wide sheet per recording: 96 columns, one row per sample,
carrying the I-VT event classification, gaze in screen pixels, pupil diameter,
per-eye validity, the presented stimulus name, and the calibration and
validation quality of that recording. Only the columns below are read.

`Recording timestamp` is in MICROSECONDS in this export, while `Recording
duration` is in milliseconds. The reader converts to milliseconds and records
the agreement between the two as a conformance check, because a silent unit
error here would rescale every duration-based feature.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator, Sequence

import numpy as np
import openpyxl

SAMPLE_COLUMNS = [
    "Recording timestamp",
    "Presented Stimulus name",
    "Eye movement type",
    "Eye movement event duration",
    "Eye movement type index",
    "Gaze point X",
    "Gaze point Y",
    "Pupil diameter left",
    "Pupil diameter right",
    "Validity left",
    "Validity right",
    "Event",
    "Event value",
]

# Constant for the whole recording; read once from the first data row.
METADATA_COLUMNS = [
    "Participant name",
    "Recording name",
    "Recording date",
    "Recording start time",
    "Recording duration",
    "Timeline name",
    "Recording Fixation filter name",
    "Recording software version",
    "Recording resolution width",
    "Recording resolution height",
    "Average calibration accuracy (degrees)",
    "Average calibration precision SD (degrees)",
    "Average calibration precision RMS (degrees)",
    "Average validation accuracy (degrees)",
    "Average validation precision SD (degrees)",
    "Average validation precision RMS (degrees)",
]

MOVEMENT_FIXATION = "Fixation"
MOVEMENT_SACCADE = "Saccade"
MOVEMENT_EYES_NOT_FOUND = "EyesNotFound"
MOVEMENT_UNCLASSIFIED = "Unclassified"

VALIDITY_VALID = "Valid"


@dataclass(frozen=True)
class TobiiStimulusSegment:
    """A contiguous run of samples showing one stimulus."""

    media: str
    index: int
    start_ms: float
    end_ms: float
    n_samples: int

    @property
    def duration_ms(self) -> float:
        return self.end_ms - self.start_ms


@dataclass
class TobiiRecording:
    """One Tobii export, reduced to the arrays the audit and features need."""

    path: Path
    metadata: dict[str, Any]
    time_ms: np.ndarray
    stimulus: np.ndarray
    movement: np.ndarray
    movement_duration_ms: np.ndarray
    movement_index: np.ndarray
    gaze_x: np.ndarray
    gaze_y: np.ndarray
    pupil_left: np.ndarray
    pupil_right: np.ndarray
    validity_left: np.ndarray
    validity_right: np.ndarray
    events: list[tuple[float, str, str]] = field(default_factory=list)
    timestamp_scale_note: str = "Recording timestamp read as microseconds, stored as milliseconds"

    @property
    def n_samples(self) -> int:
        return int(self.time_ms.size)

    @property
    def duration_ms(self) -> float:
        return float(self.time_ms[-1] - self.time_ms[0]) if self.n_samples else 0.0

    @property
    def timeline_name(self) -> str:
        return str(self.metadata.get("Timeline name", ""))

    @property
    def metadata_duration_ms(self) -> float:
        """`Recording duration` as Tobii reports it, in milliseconds."""
        return _as_float(self.metadata.get("Recording duration"))

    @property
    def duration_agreement(self) -> float:
        """Sample-span duration divided by the reported duration.

        Near 1.0 confirms the microsecond-to-millisecond conversion. A value
        near 1000 or 0.001 means the timestamp unit changed in a newer export
        and every duration feature would be wrong.
        """
        reported = self.metadata_duration_ms
        if not reported or reported != reported:
            return float("nan")
        return self.duration_ms / reported

    @property
    def sampling_rate_hz(self) -> float:
        duration_s = self.duration_ms / 1000.0
        return float(self.n_samples / duration_s) if duration_s > 0 else float("nan")

    def valid_mask(self) -> np.ndarray:
        """Samples with a usable gaze point in both eyes."""
        return (self.validity_left == VALIDITY_VALID) & (self.validity_right == VALIDITY_VALID)

    def stimulus_segments(self, exclude: Sequence[str] = ("Eyetracker Calibration",)) -> list[TobiiStimulusSegment]:
        """Contiguous runs of `Presented Stimulus name`, in presentation order.

        Calibration runs are excluded by default: they are tracker housekeeping,
        not paradigm media, and their length depends on the operator.
        """
        skip = set(exclude)
        segments: list[TobiiStimulusSegment] = []
        if self.n_samples == 0:
            return segments

        current: str | None = None
        start = 0
        end = 0
        count = 0
        index = 0

        def flush() -> None:
            nonlocal index
            if current is not None and count:
                segments.append(
                    TobiiStimulusSegment(
                        media=current,
                        index=index,
                        start_ms=float(self.time_ms[start]),
                        end_ms=float(self.time_ms[end]),
                        n_samples=count,
                    )
                )
                index += 1

        for position in range(self.n_samples):
            media = str(self.stimulus[position])
            if not media or media in skip:
                # A stray unlabelled sample does not end a presentation; the
                # instruction screens in this export are interrupted by a few
                # such samples and would otherwise split into two segments.
                continue
            if media != current:
                flush()
                current = media
                start = position
                count = 0
            end = position
            count += 1
        flush()
        return segments


def _header_index(sheet) -> dict[str, int]:
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    return {str(name): position for position, name in enumerate(header) if name is not None}


def read_tobii_recording(path: str | Path) -> TobiiRecording:
    """Stream one Tobii export into arrays, reading only the needed columns."""

    path = Path(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        columns = _header_index(sheet)
        missing = [name for name in SAMPLE_COLUMNS if name not in columns]
        if missing:
            raise ValueError(f"{path.name} is missing expected Tobii columns: {missing}")

        picks = {name: columns[name] for name in SAMPLE_COLUMNS}
        meta_picks = {name: columns[name] for name in METADATA_COLUMNS if name in columns}
        metadata: dict[str, Any] = {}
        collected: dict[str, list[Any]] = {name: [] for name in SAMPLE_COLUMNS}
        events: list[tuple[float, str, str]] = []

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not metadata:
                metadata = {name: row[position] for name, position in meta_picks.items()}
            for name, position in picks.items():
                collected[name].append(row[position])
            event = row[picks["Event"]]
            if event:
                events.append(
                    (
                        _as_float(row[picks["Recording timestamp"]]) / 1000.0,
                        str(event),
                        str(row[picks["Event value"]] or ""),
                    )
                )
    finally:
        workbook.close()

    return TobiiRecording(
        path=path,
        metadata=metadata,
        time_ms=_float_array(collected["Recording timestamp"]) / 1000.0,
        stimulus=_text_array(collected["Presented Stimulus name"]),
        movement=_text_array(collected["Eye movement type"]),
        movement_duration_ms=_float_array(collected["Eye movement event duration"]),
        movement_index=_float_array(collected["Eye movement type index"]),
        gaze_x=_float_array(collected["Gaze point X"]),
        gaze_y=_float_array(collected["Gaze point Y"]),
        pupil_left=_float_array(collected["Pupil diameter left"]),
        pupil_right=_float_array(collected["Pupil diameter right"]),
        validity_left=_text_array(collected["Validity left"]),
        validity_right=_text_array(collected["Validity right"]),
        events=events,
    )


def _as_float(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return float("nan")


def _float_array(values: Sequence[Any]) -> np.ndarray:
    return np.array([_as_float(value) for value in values], dtype=float)


def _text_array(values: Sequence[Any]) -> np.ndarray:
    return np.array(["" if value is None else str(value) for value in values], dtype=object)
