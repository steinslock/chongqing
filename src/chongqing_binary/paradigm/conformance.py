"""Check the paradigm specification against sampled raw files.

The specification states how each experiment was run. This module opens real
recordings and reports whether they agree. A disagreement is reported, never
silently coerced.
"""

from __future__ import annotations

import glob
import math
import warnings
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np

from ..paths import raw_data_root
from .spec import ParadigmSpec, load_paradigm_spec

warnings.filterwarnings("ignore")


def _status(ok: bool, partial: bool = False) -> str:
    if partial:
        return "partial"
    return "pass" if ok else "FAIL"


def check_eeg_task(spec: ParadigmSpec, task: str, raw_root: Path, n_subjects: int) -> list[dict[str, Any]]:
    import mne

    mne.set_log_level("ERROR")
    task_spec = spec.eeg(task)
    task_dir = raw_root / spec.eeg_common["raw_dir"] / task_spec.raw_dir
    subject_dirs = sorted(p for p in task_dir.glob("*") if p.is_dir())
    if not subject_dirs:
        return [{"modality": "eeg", "device": "", "task": task, "check": "task_dir_exists",
                 "expected": str(task_dir), "observed": "missing", "status": "FAIL", "n_subjects": 0}]
    step = max(1, len(subject_dirs) // n_subjects)
    sampled = subject_dirs[::step][:n_subjects]

    counts: dict[str, list[int]] = {code: [] for code in task_spec.codes}
    unexpected: Counter = Counter()
    soas: list[float] = []
    durations: list[float] = []
    mismatched_files = 0
    read_ok = 0

    contrast_codes = set(task_spec.codes)
    for subject in sampled:
        evt: list[Path] = []
        for pattern in spec.eeg_common["event_file_patterns"]:
            evt.extend(sorted(subject.glob(pattern)))
        if not evt:
            continue
        try:
            raw = mne.io.read_raw_bdf(evt[0], preload=False, verbose="ERROR")
            events, event_id = mne.events_from_annotations(raw, verbose="ERROR")
        except Exception:
            continue
        read_ok += 1
        durations.append(raw.n_times / raw.info["sfreq"])
        inv = {v: k for k, v in event_id.items()}
        observed = Counter(str(inv[code]) for code in events[:, 2])
        for code in task_spec.codes:
            counts[code].append(observed.get(code, 0))
        for code, n in observed.items():
            if code not in contrast_codes and not code.isalpha() and code not in {"Eyes Closed"}:
                unexpected[code] += n
        if task_spec.event_free and any(c.isdigit() for c in observed):
            mismatched_files += 1
        stim_codes = set(task_spec.codes_for_contrast())
        onsets = sorted(
            float(o) for o, d in zip(raw.annotations.onset, raw.annotations.description) if str(d) in stim_codes
        )
        if len(onsets) > 2:
            diffs = np.diff(onsets)
            soas.extend(diffs[diffs < 4.0].tolist())

    rows: list[dict[str, Any]] = []
    floor = float(task_spec.raw.get("min_event_file_coverage", 1.0))
    coverage = read_ok / len(sampled) if sampled else 0.0
    rows.append({"modality": "eeg", "device": "", "task": task, "check": "event_file_coverage",
                 "expected": f">= {floor:.0%}", "observed": f"{coverage:.0%} ({read_ok}/{len(sampled)} have a readable event file)",
                 "status": _status(coverage >= floor), "n_subjects": read_ok})

    for code in task_spec.codes:
        expected = task_spec.expected_total(code)
        seen = counts[code]
        if not seen:
            continue
        median = float(np.median(seen))
        present = sum(1 for x in seen if x > 0)
        presence = str(task_spec.event_codes[code].get("presence", "all_subjects"))
        required = {"all_subjects": 1.0, "most_subjects": 0.8, "optional": 0.0}.get(presence, 1.0)
        share = present / read_ok if read_ok else 0.0
        rows.append({
            "modality": "eeg", "device": "", "task": task,
            "check": f"code_{code}_present ({task_spec.condition_of(code)})",
            "expected": f"{presence} (>= {required:.0%})",
            "observed": f"median {median:.0f}/subject in {present}/{read_ok} subjects",
            "status": _status(share >= required), "n_subjects": read_ok,
        })
        if expected is not None:
            rows.append({
                "modality": "eeg", "device": "", "task": task,
                "check": f"code_{code}_count",
                "expected": f"{expected}", "observed": f"median {median:.0f}",
                "status": _status(abs(median - expected) <= max(3, expected * 0.05)),
                "n_subjects": read_ok,
            })

    if soas:
        expected_soa = task_spec.raw.get("soa_sec")
        if expected_soa is not None:
            observed_soa = float(np.median(soas))
            rows.append({"modality": "eeg", "device": "", "task": task, "check": "soa_sec",
                         "expected": f"{expected_soa}", "observed": f"{observed_soa:.3f}",
                         "status": _status(abs(observed_soa - float(expected_soa)) < 0.05),
                         "n_subjects": read_ok})

    if unexpected:
        rows.append({"modality": "eeg", "device": "", "task": task, "check": "undocumented_codes",
                     "expected": "none", "observed": str(dict(unexpected.most_common(5))),
                     "status": "FAIL", "n_subjects": read_ok})
    if mismatched_files:
        rows.append({"modality": "eeg", "device": "", "task": task, "check": "file_task_mismatch",
                     "expected": "0", "observed": f"{mismatched_files}",
                     "status": "FAIL", "n_subjects": read_ok})
    return rows


def check_fnirs_task(spec: ParadigmSpec, device: str, task: str, raw_root: Path, n_subjects: int) -> list[dict[str, Any]]:
    task_spec = spec.fnirs(device, task)
    task_dir = raw_root / task_spec.raw_dir
    if not task_dir.exists():
        return [{"modality": "fnirs", "device": device, "task": task, "check": "task_dir_exists",
                 "expected": str(task_dir), "observed": "missing", "status": "FAIL", "n_subjects": 0}]
    if device == "yiruid":
        return _check_yiruid(spec, task_spec, task_dir, n_subjects)
    return _check_bikom(spec, task_spec, task_dir, n_subjects)


def _check_yiruid(spec: ParadigmSpec, ts, task_dir: Path, n_subjects: int) -> list[dict[str, Any]]:
    from scipy.io import loadmat

    files = sorted(task_dir.rglob("*.nirs"))
    step = max(1, len(files) // n_subjects)
    sampled = files[::step][:n_subjects]
    durations, marker_counts, channel_counts, lambdas = [], [], [], []
    onsets_all: list[list[float]] = []
    for path in sampled:
        try:
            mat = loadmat(path)
        except Exception:
            continue
        t = np.asarray(mat["t"], dtype=float).reshape(-1)
        durations.append(float(t[-1] - t[0]))
        channel_counts.append(int(mat["d"].shape[1]))
        s = mat.get("s")
        idx = np.flatnonzero(np.asarray(s).reshape(-1) != 0) if s is not None and np.size(s) else np.array([], dtype=int)
        marker_counts.append(int(idx.size))
        onsets_all.append([round(float(t[i]), 1) for i in idx])
        sd = mat.get("SD")
        if sd is not None and np.size(sd):
            lambdas.append([float(x) for x in np.ravel(sd[0, 0]["Lambda"])])
    n = len(durations)
    rows = [{"modality": "fnirs", "device": ts.device, "task": ts.task, "check": "files_readable",
             "expected": f"{len(sampled)}", "observed": f"{n}", "status": _status(n == len(sampled)), "n_subjects": n}]
    if not n:
        return rows
    rows.append(_numeric_row("fnirs", ts.device, ts.task, "duration_sec", ts.duration_sec,
                             float(np.median(durations)), tol=max(3.0, ts.duration_sec * 0.03), n=n))
    rows.append(_numeric_row("fnirs", ts.device, ts.task, "n_measurements",
                             int(ts.device_raw["n_measurements"]), float(np.median(channel_counts)), tol=0, n=n))
    if lambdas:
        obs = lambdas[0]
        rows.append({"modality": "fnirs", "device": ts.device, "task": ts.task, "check": "wavelengths_nm",
                     "expected": str(ts.wavelengths_nm), "observed": str(obs),
                     "status": _status(obs == ts.wavelengths_nm), "n_subjects": n})
    expected_markers = ts.expected_marker_count
    if expected_markers is not None:
        rows.append(_numeric_row("fnirs", ts.device, ts.task, "marker_count",
                                 expected_markers, float(np.median(marker_counts)), tol=0, n=n))
    expected_onsets = ts.expected_marker_onsets_sec
    if expected_onsets and onsets_all:
        good = sum(1 for o in onsets_all
                   if len(o) == len(expected_onsets)
                   and all(abs(a - b) <= 2.0 for a, b in zip(o, expected_onsets)))
        rows.append({"modality": "fnirs", "device": ts.device, "task": ts.task, "check": "marker_onsets_sec",
                     "expected": str(expected_onsets[:6]), "observed": f"{good}/{n} within 2 s",
                     "status": _status(good == n), "n_subjects": n})
    return rows


def _check_bikom(spec: ParadigmSpec, ts, task_dir: Path, n_subjects: int) -> list[dict[str, Any]]:
    files = sorted(task_dir.rglob("*HBA_Oxy*.csv"))
    step = max(1, len(files) // n_subjects)
    sampled = files[::step][:n_subjects]
    durations, channel_counts = [], []
    sequences: Counter = Counter()
    for path in sampled:
        lines = path.read_text(encoding="utf-8-sig", errors="replace").splitlines()
        hi = next((i for i, line in enumerate(lines) if line.startswith("Probe1(")), None)
        if hi is None:
            continue
        header = [h.strip() for h in lines[hi].split(",")]
        rows_data = [line.split(",") for line in lines[hi + 1:] if line.strip()]
        channel_counts.append(sum(1 for h in header if h.startswith("CH")))
        durations.append(len(rows_data) * (1.0 / ts.sfreq_hz))
        if "Mark" in header:
            mi = header.index("Mark")
            marks = tuple(r[mi].strip() for r in rows_data if mi < len(r) and r[mi].strip() not in {"", "0"})
            sequences[marks] += 1
    n = len(durations)
    rows = [{"modality": "fnirs", "device": ts.device, "task": ts.task, "check": "files_readable",
             "expected": f"{len(sampled)}", "observed": f"{n}", "status": _status(n == len(sampled)), "n_subjects": n}]
    if not n:
        return rows
    rows.append(_numeric_row("fnirs", ts.device, ts.task, "duration_sec", ts.duration_sec,
                             float(np.median(durations)), tol=max(3.0, ts.duration_sec * 0.03), n=n))
    rows.append(_numeric_row("fnirs", ts.device, ts.task, "n_channels",
                             ts.n_channels, float(np.median(channel_counts)), tol=0, n=n))
    expected_seq = ts.raw.get("marker_sequence")
    if expected_seq is not None:
        want = tuple(expected_seq)
        optional = set(ts.device_raw.get("optional_marks", []) or [])
        required = tuple(m for m in want if m not in optional)

        def acceptable(seq: tuple[str, ...]) -> bool:
            return seq == want or tuple(m for m in seq if m not in optional) == required

        hit = sum(v for k, v in sequences.items() if acceptable(k))
        exact = sequences.get(want, 0)
        common = sequences.most_common(1)[0] if sequences else ((), 0)
        rows.append({"modality": "fnirs", "device": ts.device, "task": ts.task, "check": "marker_sequence",
                     "expected": (str(list(want)) + (f" ({sorted(optional)} optional)" if optional else "")) if want else "[] (no markers)",
                     "observed": f"{hit}/{n} acceptable ({exact} exact); most common {list(common[0])[:8]}",
                     "status": _status(hit == n), "n_subjects": n})
    elif ts.raw.get("marker_pairs"):
        want_len = int(ts.raw["marker_pairs"]) * 2 + 2
        good = sum(v for k, v in sequences.items() if len(k) == want_len)
        rows.append({"modality": "fnirs", "device": ts.device, "task": ts.task, "check": "marker_count",
                     "expected": f"{want_len} (ST + {ts.raw['marker_pairs']} A0/A1 pairs + ED)",
                     "observed": f"{good}/{n}", "status": _status(good == n), "n_subjects": n})
    return rows


def check_face(spec: ParadigmSpec, raw_root: Path) -> list[dict[str, Any]]:
    import openpyxl

    face = spec.face
    rows: list[dict[str, Any]] = []
    log_path = raw_root / face["task"]["segment_source"]
    if not log_path.exists():
        return [{"modality": "face", "device": "", "task": "task", "check": "web_log_exists",
                 "expected": str(log_path), "observed": "missing", "status": "FAIL", "n_subjects": 0}]
    wb = openpyxl.load_workbook(log_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(it)]
    records = [r for r in it if any(v is not None for v in r)]
    wb.close()

    key_col = header.index(face["task"]["key_column"])
    ids = [str(r[key_col]).strip() for r in records if r[key_col] is not None]
    rows.append(_numeric_row("face", "", "task", "subjects_in_log",
                             int(face["task"]["subjects_in_log"]), float(len(ids)), tol=0, n=len(ids)))
    rows.append({"modality": "face", "device": "", "task": "task", "check": "ids_unique",
                 "expected": f"{len(ids)}", "observed": f"{len(set(ids))}",
                 "status": _status(len(set(ids)) == len(ids)), "n_subjects": len(ids)})

    def to_sec(value: Any) -> float | None:
        if value is None:
            return None
        parts = str(value).split(":")
        if len(parts) != 4:
            return None
        try:
            h, m, s, ms = (int(x) for x in parts)
        except ValueError:
            return None
        return h * 3600 + m * 60 + s + ms / 1000.0

    complete = 0
    for movie, cfg in spec.face_movies().items():
        si = header.index(cfg["start_column"])
        ei = header.index(cfg["end_column"])
        want = float(cfg["stimulus_duration_sec"])
        tol = float(face["task"]["segment_duration_tolerance_sec"])
        observed = []
        for r in records:
            a, b = to_sec(r[si]), to_sec(r[ei])
            if a is not None and b is not None and b > a:
                observed.append(b - a)
        within = sum(1 for d in observed if abs(d - want) <= tol)
        rows.append({"modality": "face", "device": "", "task": f"movie_{movie}",
                     "check": "segment_duration_matches_stimulus",
                     "expected": f"{want:.2f} s +/- {tol:.0f} s",
                     "observed": f"median {np.median(observed):.1f} s; {within}/{len(observed)} within tolerance",
                     "status": _status(within >= 0.9 * len(observed), partial=0.5 * len(observed) <= within < 0.9 * len(observed)),
                     "n_subjects": len(observed)})
        complete = max(complete, len(observed))
    return rows


def _numeric_row(modality: str, device: str, task: str, check: str,
                 expected: float, observed: float, tol: float, n: int) -> dict[str, Any]:
    ok = abs(observed - expected) <= tol if math.isfinite(observed) else False
    return {"modality": modality, "device": device, "task": task, "check": check,
            "expected": f"{expected:g}", "observed": f"{observed:g}",
            "status": _status(ok), "n_subjects": n}


def run_conformance(spec_path: str | Path | None = None, n_subjects: int = 25,
                    raw_root: Path | None = None) -> list[dict[str, Any]]:
    spec = load_paradigm_spec(spec_path) if spec_path else load_paradigm_spec()
    root = Path(raw_root) if raw_root else raw_data_root()
    rows: list[dict[str, Any]] = []
    for task in spec.eeg_tasks():
        rows.extend(check_eeg_task(spec, task, root, n_subjects))
    for device in spec.fnirs_devices():
        for task in spec.fnirs_tasks(device):
            rows.extend(check_fnirs_task(spec, device, task, root, n_subjects))
    rows.extend(check_face(spec, root))
    return rows
